#!/usr/bin/env python3
"""協会けんぽ 都道府県別保険料額表 (Excel) を取得し、JSON に整形する。

使い方:
    python scripts/build.py [--year r08|all] [--out site] [--cache .cache]
    --year 省略時は最新年度のみ、'all' で変換可能な全年度を処理する。

出力 (--out) は年度ディレクトリごと:
    {year}/general.json   一般の被保険者 (健康保険・介護・子ども子育て支援金・厚生年金)
    {year}/optional.json  任意継続被保険者
    {year}/index.json     メタ情報 (取得元 URL, 生成時刻, 件数)
    latest/               最新年度 ({year}/ の内容と同一) のコピー
    index.json            変換できた年度一覧・最新年度
"""
from __future__ import annotations

import argparse
import io
import json
import re
import shutil
import sys
import urllib.request
import warnings
from datetime import datetime, timezone, timedelta
from pathlib import Path
from urllib.parse import urljoin

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

BASE = "https://www.kyoukaikenpo.or.jp"
INDEX_URL = f"{BASE}/about/business/insurance_rate/premium_prefectures/"
UA = "Mozilla/5.0 (compatible; jhia-tariff-bot/1.0; +https://github.com)"

PREFECTURES = [
    "北海道", "青森", "岩手", "宮城", "秋田", "山形", "福島", "茨城", "栃木", "群馬",
    "埼玉", "千葉", "東京", "神奈川", "新潟", "富山", "石川", "福井", "山梨", "長野",
    "岐阜", "静岡", "愛知", "三重", "滋賀", "京都", "大阪", "兵庫", "奈良", "和歌山",
    "鳥取", "島根", "岡山", "広島", "山口", "徳島", "香川", "愛媛", "高知", "福岡",
    "佐賀", "長崎", "熊本", "大分", "宮崎", "鹿児島", "沖縄",
]
GRADE_RE = re.compile(r"^(\d+)(?:[(（](\d+)[)）])?$")
JST = timezone(timedelta(hours=9))


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def num(v):
    # 注記テキスト(結合セル)が数値列に混入することがあるので、数値以外は欠損扱い
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    v = round(float(v), 4)
    return int(v) if v == int(v) else v


def money(v):
    return num(round(v, 2)) if isinstance(v, (int, float)) else None


# ---------------------------------------------------------------- discovery
def list_years() -> list[str]:
    html = fetch(INDEX_URL).decode("utf-8", "replace")
    years = sorted(set(re.findall(r"premium_prefectures/(r\d+)", html)))
    if not years:
        raise SystemExit("年度ページのリンクが見つかりません。ページ構造が変わった可能性があります。")
    return years


def find_xlsx_links(page_url: str) -> list[str]:
    html = fetch(page_url).decode("utf-8", "replace")
    links = re.findall(r'href="([^"]+\.xlsx)"', html, flags=re.I)
    return sorted({urljoin(page_url, l) for l in links})


def classify(wb: openpyxl.Workbook) -> str | None:
    ws = wb[wb.sheetnames[0]]
    head = " ".join(str(c) for row in ws.iter_rows(min_row=1, max_row=3, values_only=True)
                    for c in row if c)
    if "任意継続" in head:
        return "optional"
    if "健康保険・厚生年金保険の保険料額表" in head:
        return "general"
    return None


# ------------------------------------------------------------------ parsing
def strip_ws(s):
    """全角・半角の空白を除いた文字列(文字列以外はそのまま返す)。"""
    return re.sub(r"\s+", "", s) if isinstance(s, str) else s


def row_values(ws, row_idx: int) -> tuple:
    return next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))


def find_col(row: tuple, *texts: str) -> int | None:
    """row 内で空白を除いた文字列が texts のいずれかと完全一致する最初の列 index。"""
    for i, c in enumerate(row):
        if strip_ws(c) in texts:
            return i
    return None


def find_col_contains(row: tuple, *substrs: str) -> int | None:
    """row 内で substrs のいずれかを含む文字列セルの最初の列 index。"""
    for i, c in enumerate(row):
        if isinstance(c, str) and any(s in c for s in substrs):
            return i
    return None


def header_info(ws) -> tuple[int, int, str]:
    """(「等級」ヘッダ行 index(1-based), 等級列 index(0-based), 支部ラベル) を返す。

    列位置を決め打ちにせず、ヘッダーのテキスト(「等級」や支部名)を実際に探して
    対応関係を把握する。将来フォーマットが変わって列が増減・並び替わっても、
    該当するヘッダーが見つからなければ例外で分かるようにする。
    """
    pref = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
        for c in row:
            if isinstance(c, str):
                # 支部表記は年度によって「北海道支部」「北海道」の両方があり、
                # 都道府県名も古い年度では「青森県」のように正式名で書かれることが
                # ある(「北海道」はそれ自体が正式名なので対象外)。「支部」は任意に
                # しつつ、都道府県名そのものであることで他の括弧書き
                # (「（４月納付分）」等)との誤マッチを防ぐ
                m = re.search(r"[（(](.+?)(?:支部)?[）)]", c)
                if m:
                    cand = m.group(1)
                    if cand not in PREFECTURES and cand[-1:] in "都道府県" and cand[:-1] in PREFECTURES:
                        cand = cand[:-1]
                    if pref is None and cand in PREFECTURES:
                        pref = cand
        grade_col = find_col(row, "等級")
        if grade_col is not None:
            return i, grade_col, pref
    raise ValueError("「等級」ヘッダが見つかりません")


def find_amount_cols(ws, hdr: int, grade_col: int) -> tuple[int, int, int]:
    """(標準報酬月額, 報酬月額下限, 報酬月額上限) の列 index を、ヘッダーのテキスト
    (「月額」「円以上」「円未満」)を探して返す。"""
    std_col = grade_col + 1
    hdr_row = row_values(ws, hdr)
    if std_col >= len(hdr_row) or strip_ws(hdr_row[std_col]) != "月額":
        raise ValueError(f"'等級' の次の列が '月額' ヘッダではありません: {hdr_row[std_col:std_col + 1]!r}")

    min_col = max_col = None
    for offset in range(1, 4):
        row = row_values(ws, hdr + offset)
        if min_col is None:
            min_col = find_col(row, "円以上")
        if max_col is None:
            max_col = find_col(row, "円未満")
        if min_col is not None and max_col is not None:
            break
    if min_col is None or max_col is None:
        raise ValueError("報酬月額の下限/上限ヘッダ(「円以上」/「円未満」)が見つかりません")
    return std_col, min_col, max_col


def find_pair_row(ws, hdr: int) -> tuple:
    """「全額」「折半額」の見出しが並ぶ行を、hdr の直後数行の窓から探して返す。
    固定の相対位置(hdr + 1)に決め打ちせず、間に行が増えてもテキストで追従する。"""
    for offset in range(1, 4):
        row = row_values(ws, hdr + offset)
        if find_col(row, "全額") is not None and find_col(row, "折半額") is not None:
            return row
    raise ValueError("「全額」「折半額」の見出し行が見つかりません")


def classify_amount_block(ws, hdr: int, col: int) -> str | None:
    """全額/折半額の列が、ヘッダー上段のどのブロック(health/health_ltc/
    childcare_levy/pension)に属するかを、その列の見出しテキストから判別する。"""
    sub = row_values(ws, hdr - 1)[col] if hdr - 1 >= 1 else None
    title = row_values(ws, hdr - 2)[col] if hdr - 2 >= 1 else None
    text = f"{sub or ''}{title or ''}"
    if "該当しない" in text:
        return "health"
    if "該当する" in text:
        return "health_ltc"
    if "子ども" in text:
        return "childcare_levy"
    if "厚生年金" in text or "坑内員" in text or "一般" in text:
        return "pension"
    return None


def effective_text(ws) -> str:
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        for c in row:
            if isinstance(c, str) and ("分" in c) and ("令和" in c or "平成" in c) and "から" in c or \
               isinstance(c, str) and re.match(r"^[（(]令和.*分.*[）)]$", c.strip()):
                return c.strip()
    return ""


def parse_general(wb) -> list[dict]:
    out = []
    for code, (name, sn) in enumerate(zip(PREFECTURES, wb.sheetnames), 1):
        ws = wb[sn]
        hdr, grade_col, label = header_info(ws)
        if label != name:
            raise ValueError(f"シート {sn!r}: 支部ラベル {label!r} が想定 {name!r} と不一致")
        std_col, min_col, max_col = find_amount_cols(ws, hdr, grade_col)

        # 「全額」「折半額」ヘッダーの列を左から順に拾い、同じ順序で対になっていると
        # みなす(年度によってブロック数(3 or 4)が違うため、個数を決め打ちしない)。
        pair_row = find_pair_row(ws, hdr)
        total_cols = [i for i, c in enumerate(pair_row) if strip_ws(c) == "全額"]
        half_cols = [i for i, c in enumerate(pair_row) if strip_ws(c) == "折半額"]
        if len(total_cols) != len(half_cols):
            raise ValueError(f"シート {sn!r}: 「全額」({len(total_cols)})と「折半額」({len(half_cols)})の列数が不一致")

        blocks: dict[str, tuple[int, int]] = {}
        for tcol, hcol in zip(total_cols, half_cols):
            block = classify_amount_block(ws, hdr, tcol)
            if block is None:
                raise ValueError(f"シート {sn!r}: {tcol} 列目の金額ブロック(健康保険/介護/支援金/年金)を判別できません")
            if block in blocks:
                raise ValueError(f"シート {sn!r}: 金額ブロック {block!r} が複数見つかりました")
            blocks[block] = (tcol, hcol)

        missing = {"health", "health_ltc", "pension"} - set(blocks)
        if missing:
            raise ValueError(f"シート {sn!r}: 必須の金額ブロックが見つかりません: {missing}")

        hdr_row = row_values(ws, hdr)
        rates = {block: num(hdr_row[tcol]) for block, (tcol, _) in blocks.items()}

        for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
            g = row[grade_col]
            m = GRADE_RE.match(str(g).strip()) if g is not None else None
            if not m or row[std_col] is None:
                continue

            def amount(block):
                if block not in blocks:
                    return None, None
                tcol, hcol = blocks[block]
                return row[tcol], row[hcol]

            health_full, health_half = amount("health")
            health_ltc_full, health_ltc_half = amount("health_ltc")
            childcare_levy_full, childcare_levy_half = amount("childcare_levy")
            pension_full, pension_half = amount("pension")
            out.append({
                "prefecture_code": code,
                "prefecture_name": name,
                "grade": int(m.group(1)),
                "pension_grade": int(m.group(2)) if m.group(2) else None,
                "remuneration_std": num(row[std_col]),
                "remuneration_min": num(row[min_col]),
                "remuneration_max": num(row[max_col]),
                "rate_health": rates.get("health"),
                "rate_health_ltc": rates.get("health_ltc"),
                "rate_childcare_levy": rates.get("childcare_levy"),
                "rate_pension": rates.get("pension"),
                "health_full": money(health_full),
                "health_half": money(health_half),
                "health_ltc_full": money(health_ltc_full),
                "health_ltc_half": money(health_ltc_half),
                "childcare_levy_full": money(childcare_levy_full),
                "childcare_levy_half": money(childcare_levy_half),
                "pension_full": money(pension_full),
                "pension_half": money(pension_half),
            })
    return out


def parse_optional(wb) -> list[dict]:
    out = []
    for code, (name, sn) in enumerate(zip(PREFECTURES, wb.sheetnames), 1):
        ws = wb[sn]
        if sn != f"{code:02d}{name}":
            raise ValueError(f"シート名 {sn!r} が想定 {code:02d}{name!r} と不一致")
        hdr, grade_col, label = header_info(ws)
        if label != name:
            raise ValueError(f"シート {sn!r}: 支部ラベル {label!r} が想定 {name!r} と不一致")
        std_col, min_col, max_col = find_amount_cols(ws, hdr, grade_col)

        sub_row = row_values(ws, hdr - 1)
        plain_col = find_col_contains(sub_row, "該当しない")
        care_col = find_col_contains(sub_row, "該当する")
        if plain_col is None or care_col is None:
            raise ValueError(f"シート {sn!r}: 介護保険の該当有無ヘッダが見つかりません")

        hdr_row = row_values(ws, hdr)
        r_plain, r_care = hdr_row[plain_col], hdr_row[care_col]
        for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
            g = row[grade_col]
            if not isinstance(g, int) or row[std_col] is None:
                continue
            out.append({
                "prefecture_code": code,
                "prefecture_name": name,
                "grade": g,
                "remuneration_std": num(row[std_col]),
                "remuneration_min": num(row[min_col]),
                "remuneration_max": num(row[max_col]),
                "rate_health": num(r_plain),
                "rate_health_ltc": num(r_care),
                "health_full": money(row[plain_col]),
                "health_ltc_full": money(row[care_col]),
            })
    return out


PARSERS = {"general": parse_general, "optional": parse_optional}

# *.ja.json 用の日本語キー (英語キー → 日本語キー。出力順も兼ねる)
JA_KEYS = {
    "general": {
        "prefecture_code": "都道府県コード",
        "prefecture_name": "都道府県",
        "grade": "健保等級",
        "pension_grade": "厚年等級",
        "remuneration_std": "標準報酬月額",
        "remuneration_min": "報酬月額下限",
        "remuneration_max": "報酬月額上限",
        "rate_health": "健康保険料率",
        "rate_health_ltc": "介護保険料率込み健康保険料率",
        "rate_childcare_levy": "子ども・子育て支援金率",
        "rate_pension": "厚生年金保険料率",
        "health_full": "健康保険料(全額)",
        "health_half": "健康保険料(折半額)",
        "health_ltc_full": "健康保険料・介護あり(全額)",
        "health_ltc_half": "健康保険料・介護あり(折半額)",
        "childcare_levy_full": "子ども・子育て支援金(全額)",
        "childcare_levy_half": "子ども・子育て支援金(折半額)",
        "pension_full": "厚生年金保険料(全額)",
        "pension_half": "厚生年金保険料(折半額)",
    },
    "optional": {
        "prefecture_code": "都道府県コード",
        "prefecture_name": "都道府県",
        "grade": "等級",
        "remuneration_std": "標準報酬月額",
        "remuneration_min": "報酬月額下限",
        "remuneration_max": "報酬月額上限",
        "rate_health": "保険料率(介護なし)",
        "rate_health_ltc": "保険料率(介護あり)",
        "health_full": "保険料(介護なし)",
        "health_ltc_full": "保険料(介護あり)",
    },
}


def to_ja(kind: str, rows: list[dict]) -> list[dict]:
    keys = JA_KEYS[kind]
    if set(rows[0]) != set(keys):
        raise ValueError(f"{kind}: 日本語キー対応表が不一致: {set(rows[0]) ^ set(keys)}")
    return [{ja: r[en] for en, ja in keys.items()} for r in rows]


def validate(kind: str, rows: list[dict]) -> None:
    if len({r["prefecture_code"] for r in rows}) != 47:
        raise ValueError(f"{kind}: 47 都道府県そろっていません")
    per = {}
    for r in rows:
        per.setdefault(r["prefecture_code"], []).append(r["grade"])
    counts = {len(v) for v in per.values()}
    if len(counts) != 1:
        raise ValueError(f"{kind}: 都道府県ごとの等級数が不一致: {counts}")
    for code, grades in per.items():
        if grades != sorted(grades) or len(set(grades)) != len(grades):
            raise ValueError(f"{kind}: 等級が単調増加でない (pref={code})")
    prem = "health_full"
    if any(r[prem] is None or r[prem] <= 0 for r in rows):
        raise ValueError(f"{kind}: 保険料が欠損/非正の行があります")


def build_year(year: str, out: Path, cache: Path) -> dict:
    """1 年度分を取得・変換し、out 直下 (呼び出し側が年度ディレクトリを渡す) に書き出す。"""
    page = f"{INDEX_URL}{year}"
    print(f"年度ページ: {page}")
    out.mkdir(parents=True, exist_ok=True)
    cache.mkdir(parents=True, exist_ok=True)

    found: dict[str, dict] = {}
    for url in find_xlsx_links(page):
        data = fetch(url)
        (cache / Path(url).name).write_bytes(data)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        kind = classify(wb)
        print(f"  {url} -> {kind} ({len(wb.sheetnames)} sheets)")
        if kind and kind not in found:
            found[kind] = {"wb": wb, "url": url}
    missing = set(PARSERS) - set(found)
    if missing:
        raise SystemExit(f"[{year}] Excel が見つかりません: {missing}")

    files = {}
    for kind, parse in PARSERS.items():
        wb = found[kind]["wb"]
        rows = parse(wb)
        validate(kind, rows)
        ws = wb[wb.sheetnames[0]]
        title = " ".join(str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if c)
        (out / f"{kind}.json").write_text(
            json.dumps(rows, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        (out / f"{kind}.ja.json").write_text(
            json.dumps(to_ja(kind, rows), ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        files[kind] = {"file": f"{kind}.json", "file_ja": f"{kind}.ja.json",
                       "rows": len(rows), "title": title, "source": found[kind]["url"]}
        print(f"  {kind}.json, {kind}.ja.json: {len(rows)} rows")

    index = {
        "generated_at": datetime.now(JST).isoformat(timespec="seconds"),
        "year_page": page,
        "year": year,
        "files": files,
    }
    (out / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return index


def copy_as_latest(year_dir: Path, latest_dir: Path) -> None:
    latest_dir.mkdir(parents=True, exist_ok=True)
    names = ["general.json", "general.ja.json", "optional.json", "optional.ja.json", "index.json"]
    for name in names:
        (latest_dir / name).write_bytes((year_dir / name).read_bytes())


# --------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", help="年度 (例: r08)。'all' で全年度。省略時は最新のみ")
    ap.add_argument("--out", default="site")
    ap.add_argument("--cache", default=".cache")
    a = ap.parse_args()

    years = list_years()
    latest = years[-1]
    if a.year is None:
        targets = [latest]
    elif a.year == "all":
        targets = years
    elif a.year in years:
        targets = [a.year]
    else:
        raise SystemExit(f"年度 {a.year} は存在しません (候補: {years})")

    out, cache = Path(a.out), Path(a.cache)
    built = []
    for year in targets:
        try:
            build_year(year, out / year, cache / year)
            built.append(year)
        except (Exception, SystemExit) as e:
            # 単一年度を明示指定した場合は想定通りの失敗として即座に止める。
            # --year all の場合は、レイアウトが異なる古い年度のために全体を
            # 止めないよう、警告を出してスキップし、残りの年度は処理を続ける。
            if len(targets) == 1:
                raise
            print(f"[警告] {year}: スキップ ({e})", file=sys.stderr)
            shutil.rmtree(out / year, ignore_errors=True)

    if not built:
        raise SystemExit("どの年度も変換できませんでした。")

    if latest in built:
        copy_as_latest(out / latest, out / "latest")

    (out / "index.json").write_text(json.dumps({
        "generated_at": datetime.now(JST).isoformat(timespec="seconds"),
        "years": built,
        "latest": latest if latest in built else None,
    }, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
